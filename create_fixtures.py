"""Generate all required test fixtures for ClauseLens Trust."""

import os
import fitz
import docx

fixtures_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'tests', 'fixtures')
expected_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'tests', 'expected')
os.makedirs(fixtures_dir, exist_ok=True)
os.makedirs(expected_dir, exist_ok=True)

# 1. Standard 4-conflict sample freelance contract (PDF)
# Conflicts:
# - Payment within sixty (60) days of invoice (>30 days limit -> Conflict, High severity)
# - Automatically renew for successive twelve (12) month terms unless 90 days notice (Auto-renewal disallowed -> Conflict, High notice)
# - Liability shall be unlimited (Unlimited disallowed -> Conflict, High severity)
# - Exclusive jurisdiction in the courts of Singapore (Singapore not in India -> Conflict, Medium severity)
# Also termination for convenience (30 days notice)
doc1 = fitz.open()
page1 = doc1.new_page()
text1 = """SAMPLE FREELANCE SERVICE AGREEMENT

1. Services and Deliverables
The Consultant shall perform custom software development and consulting services as requested by the Client.

2. Fees and Payment Terms
All invoices submitted by Consultant shall be due and payment within sixty (60) days of invoice.

3. Term and Renewal
This Agreement shall commence on the Effective Date and shall automatically renew for successive twelve (12) month terms unless 90 days notice of cancellation is provided in writing by either party prior to expiration.

4. Limitation of Liability
Notwithstanding anything to the contrary herein, Consultant's liability shall be unlimited for any and all claims, losses, or damages arising out of or related to this Agreement.

5. Termination
Either party may terminate this Agreement for convenience upon thirty (30) days prior written notice.

6. Governing Law and Jurisdiction
This Agreement shall be governed by and construed in accordance with the laws of Singapore, and the parties hereby submit to the exclusive jurisdiction of the courts of Singapore.
"""
page1.insert_textbox(fitz.Rect(50, 50, 550, 750), text1, fontsize=10, fontname="helv")
pdf1_path = os.path.join(fixtures_dir, 'sample_freelance_contract.pdf')
doc1.save(pdf1_path)
doc1.close()
print(f"Created: {pdf1_path}")

# 2. Ambiguous payment contract (PDF)
doc2 = fitz.open()
page2 = doc2.new_page()
text2 = """STANDARD CONSULTING AGREEMENT

Section 1. Engagement
The Service Provider agrees to deliver technical advisory services.

Section 2. Compensation
The Client shall pay all undisputed amounts within a reasonable time following submission of an itemized statement.

Section 3. Term
This Agreement terminates upon completion of the deliverables unless extended by mutual agreement.

Section 4. Liability
The aggregate liability of either party shall not exceed the total fees paid under this Agreement.

Section 5. Governing Law
This Agreement shall be governed by the laws of India and subject to the jurisdiction of courts in Chennai.
"""
page2.insert_textbox(fitz.Rect(50, 50, 550, 750), text2, fontsize=10, fontname="helv")
pdf2_path = os.path.join(fixtures_dir, 'ambiguous_payment_contract.pdf')
doc2.save(pdf2_path)
doc2.close()
print(f"Created: {pdf2_path}")

# 3. Clean compliant contract (PDF)
doc3 = fitz.open()
page3 = doc3.new_page()
text3 = """MASTER SERVICES AGREEMENT

1. Scope of Work
Vendor shall provide design and engineering services as described in attached Statement of Work.

2. Invoicing and Payment
Client agrees to remit payment within 15 calendar days of receipt of each valid invoice.

3. Term
The term of this Agreement is twelve (12) months from the Effective Date, with no automatic renewal.

4. Limitation of Liability
The total aggregate liability of Vendor arising under this Agreement shall not exceed the total fees paid by Client during the preceding three months.

5. Termination
Either party may terminate this Agreement with thirty (30) days prior written notice for convenience.

6. Jurisdiction
This Agreement is governed by the laws of India, and the courts of Mumbai shall have exclusive jurisdiction.
"""
page3.insert_textbox(fitz.Rect(50, 50, 550, 750), text3, fontsize=10, fontname="helv")
pdf3_path = os.path.join(fixtures_dir, 'clean_compliant_contract.pdf')
doc3.save(pdf3_path)
doc3.close()
print(f"Created: {pdf3_path}")

# 4. Vendor Agreement (DOCX)
# Net 45 days, 24 months renewal, unlimited liability, London courts
doc4 = docx.Document()
doc4.add_heading('VENDOR SUPPLY AGREEMENT', 0)
doc4.add_paragraph('This Vendor Supply Agreement is entered into between Buyer and Vendor.')
doc4.add_heading('1. Payment Terms', level=1)
doc4.add_paragraph('Buyer shall make payment within forty-five (45) days of delivery of invoice.')
doc4.add_heading('2. Term and Renewal', level=1)
doc4.add_paragraph('This contract shall continue for an additional 24 month renewal period unless terminated with sixty (60) days notice.')
doc4.add_heading('3. Liability and Indemnity', level=1)
doc4.add_paragraph('Vendor shall bear uncapped liability and indemnifies Buyer against all losses arising from breach.')
doc4.add_heading('4. Jurisdiction', level=1)
doc4.add_paragraph('This agreement shall be governed by the laws of England, and the courts of London shall have exclusive jurisdiction.')
docx4_path = os.path.join(fixtures_dir, 'vendor_agreement.docx')
doc4.save(docx4_path)
print(f"Created: {docx4_path}")

# 5. Consulting Contract with Table (DOCX)
doc5 = docx.Document()
doc5.add_heading('PROFESSIONAL SERVICES CONTRACT', 0)
doc5.add_paragraph('Agreement between Client and Consultant.')
doc5.add_heading('1. Payment Schedule', level=1)
table = doc5.add_table(rows=2, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Milestone'
hdr_cells[1].text = 'Deliverable'
hdr_cells[2].text = 'Payment Terms'
row_cells = table.rows[1].cells
row_cells[0].text = 'M1'
row_cells[1].text = 'Initial Architecture'
row_cells[2].text = 'Advance payment required prior to commencement.'

doc5.add_heading('2. Termination', level=1)
doc5.add_paragraph('Either party may terminate this agreement upon sixty (60) days written notice.')
doc5.add_heading('3. Liability', level=1)
doc5.add_paragraph('Total liability of either party shall be limited to the fees actually paid under this contract.')
doc5.add_heading('4. Governing Law', level=1)
doc5.add_paragraph('This contract shall be governed by the laws of India, and disputes shall be resolved in Chennai.')
docx5_path = os.path.join(fixtures_dir, 'consulting_contract.docx')
doc5.save(docx5_path)
print(f"Created: {docx5_path}")

# 6. Scanned or empty PDF (<50 chars)
doc6 = fitz.open()
page6 = doc6.new_page()
page6.insert_text((50, 50), "Scanned doc", fontsize=8) # Only 11 chars
pdf6_path = os.path.join(fixtures_dir, 'scanned_empty.pdf')
doc6.save(pdf6_path)
doc6.close()
print(f"Created: {pdf6_path}")

# 7. Encrypted PDF
doc7 = fitz.open()
page7 = doc7.new_page()
page7.insert_text((50, 50), "Confidential agreement text.", fontsize=10)
pdf7_path = os.path.join(fixtures_dir, 'encrypted.pdf')
doc7.save(pdf7_path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret123", owner_pw="owner123")
doc7.close()
print(f"Created: {pdf7_path}")

# 8. Corrupt file
corrupt_path = os.path.join(fixtures_dir, 'corrupt.pdf')
with open(corrupt_path, 'wb') as f:
    f.write(b'%PDF-1.4\n%not-a-valid-pdf-body-garbage-bytes\x00\xff\xfe\x00')
print(f"Created: {corrupt_path}")

# 9. Wrong-type file
wrong_path = os.path.join(fixtures_dir, 'invalid_contract.jpg')
with open(wrong_path, 'wb') as f:
    f.write(b'\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x01\x00`\x00`\x00\x00')
print(f"Created: {wrong_path}")

print("All fixtures successfully generated!")
