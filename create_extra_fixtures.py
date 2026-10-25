"""Create extra demo contracts (PDF and DOCX)."""

import os
import fitz
import docx

fixtures_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'tests', 'fixtures')

# 1. Strict Non-Disclosure Agreement (PDF)
doc_nda = fitz.open()
page_nda = doc_nda.new_page()
nda_text = """MUTUAL NON-DISCLOSURE AGREEMENT

1. Purpose and Scope
The parties wish to explore a business relationship and disclose certain confidential information.

2. Payment and Costs
Each party shall bear its own expenses. In the event of permitted consulting fee disbursements, payment within 60 days of invoice shall apply.

3. Term and Renewal
This Agreement shall remain in effect for two (2) years and shall automatically renew for additional 12 month terms unless written notice is given sixty (60) days prior.

4. Indemnity and Liability
Recipient agrees to indemnify and hold harmless the Discloser against all damages, liabilities, and losses without limitation.

5. Termination
Either party may terminate this Agreement upon thirty (30) days prior written notice.

6. Governing Law
This Agreement shall be governed exclusively by the laws of California, USA, and the courts of California shall have exclusive jurisdiction.
"""
page_nda.insert_textbox(fitz.Rect(50, 50, 550, 750), nda_text, fontsize=10, fontname="helv")
nda_path = os.path.join(fixtures_dir, 'nda_strict_confidentiality.pdf')
doc_nda.save(nda_path)
doc_nda.close()
print(f"Created: {nda_path}")

# 2. Software Services Agreement (DOCX)
doc_saas = docx.Document()
doc_saas.add_heading('SOFTWARE SUBSCRIPTION AGREEMENT', 0)
doc_saas.add_paragraph('This Software Subscription Agreement is entered into between Provider and Subscriber.')

doc_saas.add_heading('1. Fees and Invoicing', level=1)
doc_saas.add_paragraph('Subscription fees shall be payable within sixty (60) days of receipt of invoice.')

doc_saas.add_heading('2. Subscription Term & Auto-Renewal', level=1)
doc_saas.add_paragraph('The initial subscription term is 12 months. This agreement will automatically renew for successive twelve (12) month terms unless ninety (90) days notice of non-renewal is provided.')

doc_saas.add_heading('3. Limitation of Liability', level=1)
doc_saas.add_paragraph('Neither party shall have unlimited liability. Provider total aggregate liability shall not exceed the subscription fees paid in the prior twelve months.')

doc_saas.add_heading('4. Termination for Convenience', level=1)
doc_saas.add_paragraph('Provider may terminate this agreement at will upon fifteen (15) days written notice.')

doc_saas.add_heading('5. Dispute Resolution and Jurisdiction', level=1)
doc_saas.add_paragraph('This agreement is governed by the laws of Singapore, and all disputes shall be resolved in Singapore courts.')

saas_path = os.path.join(fixtures_dir, 'software_license_agreement.docx')
doc_saas.save(saas_path)
print(f"Created: {saas_path}")
